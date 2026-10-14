"""render_docx_preview_html: HTML-предпросмотр .docx без сохранения третьего
файла (см. npa_pipeline.docx_preview — почему это нужно и чем не является
точной копией исходной вёрстки)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from npa_pipeline.docx_preview import render_docx_preview_html


def _build_docx(path: Path) -> None:
    doc = Document()
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p1.add_run("ЗАГОЛОВОК")
    run.bold = True

    doc.add_paragraph("Обычный текст статьи.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Строка 1"
    table.cell(0, 1).text = "Значение 1"
    table.cell(1, 0).text = "Строка 2"
    table.cell(1, 1).text = "Значение 2"

    doc.save(str(path))


def test_renders_paragraphs_in_order_with_alignment_and_bold(tmp_path: Path):
    path = tmp_path / "doc.docx"
    _build_docx(path)
    html = render_docx_preview_html(path)

    assert html.startswith("<!doctype html>")
    assert "<strong>ЗАГОЛОВОК</strong>" in html
    assert 'text-align:center' in html
    assert "Обычный текст статьи." in html
    # заголовок (центр, жирный) должен идти раньше обычного абзаца в выдаче
    assert html.index("ЗАГОЛОВОК") < html.index("Обычный текст статьи.")


def test_renders_table_as_html_table(tmp_path: Path):
    path = tmp_path / "doc.docx"
    _build_docx(path)
    html = render_docx_preview_html(path)

    assert "<table>" in html
    assert "Строка 1" in html
    assert "Значение 2" in html


def test_merged_cell_not_duplicated_in_row(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Объединённая ячейка"
    path = tmp_path / "merged.docx"
    doc.save(str(path))

    html = render_docx_preview_html(path)
    assert html.count("Объединённая ячейка") == 1


def test_raises_for_non_docx_file(tmp_path: Path):
    path = tmp_path / "not_really.docx"
    path.write_bytes(b"this is not a zip/docx file at all")
    with pytest.raises(Exception):
        render_docx_preview_html(path)
